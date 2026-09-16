"""Alberta Table 93: Non-Board Feed Grain Prices (XLSX).

Terminal / local elevator prices at named Alberta delivery points, already
$ per tonne. These are NOT farm-gate prices, which is why the registry types this
source ``terminal_local_price``: it must never be averaged into an
``average_farm_price`` series.

Layout, verified against the downloaded workbook:

    row 1     title
    row 3     market names, merged across their crop columns
              ('Lethbridge', 'Calgary', 'Red Deer', ...)
    row 4     crop names under each market ('Wheat', 'Oats', 'Barley', ...)
    row 6     the unit, printed once: '$/tonne'
    row 8+    data: col A = marketing year ('2003-04......')

The market name lives on a *different row* from the crop name and is merged across
several columns, so both rows are forward-filled before pairing. Reading only row 4
would attach every series to a market of None; reading only row 3 would do the
reverse. Neither the market nor the crop may be guessed from column position,
because the set of markets and the crops carried at each one both vary.

A single '-' cell means "no quote that year", never zero.
"""

from __future__ import annotations

import re
from typing import Any, Iterator

import openpyxl

from .. import crops, units
from ..models import Observation, make_observation_id
from .base import ParseContext, raw_repr_of

SHEET = "Table 93"

MARKET_ROW = 3
CROP_ROW = 4
UNIT_ROW = 6
FIRST_DATA_ROW = 8
COL_YEAR = 1
FIRST_DATA_COL = 3        # column B is a spacer

# Marketing years look like '2003-04......' -- the dots are leader padding.
MKT_YEAR_RE = re.compile(r"^(\d{4})-(\d{2})")

BLANK_TOKENS = {"-", "--", "---", "...", "\u2026", "n/a", "na", " ", ""}

# Header/annotation cells that must not be mistaken for a market or a crop.
_NOT_SERIES = re.compile(
    r"(^\s*\$|^\s*source|^\s*symbols|^\s*$/tonne|non-board|price\s+list)", re.I,
)


def _blank(v: Any) -> bool:
    if v is None:
        return True
    return str(v).strip().lower() in BLANK_TOKENS


def _text(v: Any) -> str | None:
    if _blank(v) or not isinstance(v, str):
        return None
    t = re.sub(r"\s+", " ", v).strip()
    if not t or _NOT_SERIES.search(t):
        return None
    return t


def _to_float(v: Any) -> float | None:
    if _blank(v) or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    text = re.sub(r"[^\d.\-]", "", str(v))
    try:
        return float(text)
    except ValueError:
        return None


def _forward_fill(values: list[Any]) -> list[str | None]:
    """Carry the last non-empty text rightward, for merged/grouped headers.

    The market row stores a name once and leaves the rest of its span blank, so
    forward-filling both header rows yields the (market, crop) pair that each data
    column belongs to.
    """
    out: list[str | None] = []
    last: str | None = None
    for v in values:
        t = _text(v)
        if t is not None:
            last = t
        out.append(last)
    return out


def _marketing_year(v: Any) -> str | None:
    if v is None:
        return None
    m = MKT_YEAR_RE.match(str(v).strip())
    return f"{m.group(1)}-{m.group(2)}" if m else None


def header_pairs(ws) -> dict[int, tuple[str, str]]:
    """{column index -> (market, crop)} for every priced column on the sheet."""
    width = ws.max_column
    markets = _forward_fill([ws.cell(MARKET_ROW, c).value for c in range(1, width + 1)])
    crop_names = _forward_fill([ws.cell(CROP_ROW, c).value for c in range(1, width + 1)])

    pairs: dict[int, tuple[str, str]] = {}
    for c in range(FIRST_DATA_COL, width + 1):
        market, crop_name = markets[c - 1], crop_names[c - 1]
        if not market or not crop_name:
            continue
        pairs[c] = (market, crop_name)
    return pairs


def _sheet_unit(ws) -> str:
    """The unit printed once on the sheet, e.g. '$/tonne'."""
    for c in range(1, ws.max_column + 1):
        v = ws.cell(UNIT_ROW, c).value
        if isinstance(v, str) and "$" in v:
            return units.normalize_unit(v)
    return ""
def iter_rows(ctx: ParseContext, options: dict[str, Any]) -> Iterator[Observation]:
    """Yield terminal/local feed-grain prices from Table 93."""
    src = ctx.source
    wb = openpyxl.load_workbook(src.raw_path, data_only=True)
    try:
        ws = wb[options.get("sheet") or SHEET]
        unit = _sheet_unit(ws)
        pairs = header_pairs(ws)

        for r in range(FIRST_DATA_ROW, ws.max_row + 1):
            year = _marketing_year(ws.cell(r, COL_YEAR).value)
            if year is None:
                continue
            for c, (market, crop_name) in sorted(pairs.items()):
                value = _to_float(ws.cell(r, c).value)
                if value is None:
                    continue
                yield _observation(ctx, crop_name, market, value, unit, year, r, c)
    finally:
        wb.close()


def _observation(ctx: ParseContext, crop_label: str, market: str, value: float,
                 unit: str, year: str, row_no: int, col_no: int) -> Observation:
    src = ctx.source
    clean = crops.clean_label(crop_label)
    conv = units.to_cad_per_tonne(value, unit=unit or "tonne", crop_key=clean)
    return Observation(
        observation_id=make_observation_id(src.source_id, clean, market, year,
                                           f"r{row_no}c{col_no}"),
        crop_id=crops.crop_for(clean),
        source_id=src.source_id,
        source_title=src.title,
        publisher=src.publisher,
        price_type=src.price_type,
        region="Alberta",
        reference_date=year,
        date_granularity="marketing-year",
        original_value=value,
        original_unit=unit or "tonne",
        currency="CAD",
        normalized_value=round(conv.cad_per_tonne, 4) if conv.ok else None,
        normalized_unit="CAD/tonne",
        conversion_basis=conv.basis,
        conversion_factor=conv.factor,
        conversion_detail=conv.detail,
        record_origin="observed",
        source_commodity=crop_label,
        variant=crops.variant_for(clean),
        market=market,
        doc_file=src.raw_name,
        doc_table=SHEET,
        source_url=src.download_url,
        licence="Open Government Licence - Alberta",
        retrieved_at=ctx.retrieved_at,
        raw_sha256=ctx.sha256,
        raw_repr=raw_repr_of({"market": market, "crop": crop_label, "year": year,
                              "value": value, "unit": unit or "tonne",
                              "row": row_no, "col": col_no}),
    )
