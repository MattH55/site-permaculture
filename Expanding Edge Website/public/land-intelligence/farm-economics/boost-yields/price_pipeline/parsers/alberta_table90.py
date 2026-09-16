"""Alberta Table 90: Special Crops -- area, yield, production and price (XLSX).

Static 2005-2014 series. The workbook declares the price unit *per row* in its own
column, so the unit is read from the data rather than assumed for the sheet. That
matters because this one table prices corn per bushel and potatoes per
hundredweight; a sheet-wide unit would convert one of them into fiction.

Layout, verified against the downloaded workbook:

    row 1     title
    rows 3-5  header band: 'Harvested Area'/Acres, 'Yield'/Per Acre,
              'Production'/Tonnes, 'Average Value'/'$/Unit'
    col A     year, e.g. '2005......' or '2014p'   (trailing p = preliminary)
    col B     area, OR a lone crop name introducing a block
    col D     yield (col E = that yield's unit)
    col F     production, tonnes
    col H     average value (col I = the price unit, e.g. '/bu.' or '/cwt.')

Crop blocks are opened by a lone name in column B (Corn for Grain, Dry Peas, Dry
Beans, Lentils, Mustard Seed, Triticale, Potatoes, Fodder Corn). A block ends at
the next name, so each data row is attributed to the block it sits inside -- by
position, never by proximity.

A cell holding only a dash or dots means "not published", never zero, so it is
skipped rather than emitted as 0.
"""

from __future__ import annotations

import re
from typing import Any, Iterator

import openpyxl

from .. import crops, units
from ..models import Observation, make_observation_id
from .base import ParseContext, raw_repr_of

SHEET = "Table 90"

YEAR_RE = re.compile(r"^(\d{4})\s*([prx]?)")
# 1-based column indices (openpyxl convention).
COL_YEAR = 1
COL_LABEL = 2
COL_PRICE = 8
COL_PRICE_UNIT = 9

BLANK_TOKENS = {"-", "--", "---", "...", "\u2026", "n/a", "na", ""}

# Header words that share column B with the crop names but are not crop names.
_NOT_A_CROP = re.compile(
    r"(harvested\s+area|yield|production|average\s+value|acres|per\s+acre"
    r"|tonnes|\$/unit|^source:|^symbols)", re.I,
)


def _blank(v: Any) -> bool:
    if v is None:
        return True
    return str(v).strip().lower() in BLANK_TOKENS


def _to_float(v: Any) -> float | None:
    if _blank(v) or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    text = re.sub(r"[^\d.\-]", "", str(v))
    try:
        return float(text)
    except ValueError:
        return None


def _year_status(v: Any) -> tuple[str | None, str | None]:
    """(year, status) from a year cell like '2014p' or '2009r'."""
    if v is None:
        return None, None
    m = YEAR_RE.match(str(v).strip())
    if not m:
        return None, None
    year, suffix = m.group(1), m.group(2)
    return year, (suffix or None)


def _looks_like_crop(text: str) -> bool:
    if _NOT_A_CROP.search(text):
        return False
    return not _blank(text)


def iter_rows(ctx: ParseContext, options: dict[str, Any]) -> Iterator[Observation]:
    """Yield average farm prices from Table 90."""
    src = ctx.source
    wb = openpyxl.load_workbook(src.raw_path, data_only=True)
    try:
        ws = wb[options.get("sheet") or SHEET]
        current_crop: str | None = None
        for r in range(1, ws.max_row + 1):
            label_cell = ws.cell(r, COL_LABEL).value
            year, status = _year_status(ws.cell(r, COL_YEAR).value)

            # No year in column A means this is not a data row. If it carries a
            # real crop name it opens a new block; otherwise it is header prose.
            if year is None:
                if isinstance(label_cell, str) and _looks_like_crop(label_cell.strip()):
                    current_crop = label_cell.strip()
                continue

            if current_crop is None:
                continue

            value = _to_float(ws.cell(r, COL_PRICE).value)
            if value is None:
                continue
            raw_unit = ws.cell(r, COL_PRICE_UNIT).value
            unit = units.normalize_unit(str(raw_unit) if raw_unit else None)
            if not unit:
                # No declared unit means no honest conversion is possible.
                continue

            yield _observation(ctx, current_crop, value, unit, year, status, r)
    finally:
        wb.close()


def _observation(ctx: ParseContext, crop_label: str, value: float, unit: str,
                 year: str, status: str | None, row_no: int) -> Observation:
    src = ctx.source
    clean = crops.clean_label(crop_label)
    conv = units.to_cad_per_tonne(value, unit=unit, crop_key=clean)
    return Observation(
        observation_id=make_observation_id(src.source_id, clean, year, unit,
                                           f"r{row_no}"),
        crop_id=crops.crop_for(clean),
        source_id=src.source_id,
        source_title=src.title,
        publisher=src.publisher,
        price_type=src.price_type,
        region="Alberta",
        reference_date=year,
        date_granularity="year",
        original_value=value,
        original_unit=unit,
        currency="CAD",
        normalized_value=round(conv.cad_per_tonne, 4) if conv.ok else None,
        normalized_unit="CAD/tonne",
        conversion_basis=conv.basis,
        conversion_factor=conv.factor,
        conversion_detail=conv.detail,
        record_origin="observed",
        source_commodity=crop_label,
        variant=crops.variant_for(clean),
        year_status=status,
        doc_file=src.raw_name,
        doc_table=SHEET,
        source_url=src.download_url,
        licence="Open Government Licence - Alberta",
        retrieved_at=ctx.retrieved_at,
        raw_sha256=ctx.sha256,
        raw_repr=raw_repr_of({"crop": crop_label, "year": year, "value": value,
                              "unit": unit, "row": row_no}),
    )
